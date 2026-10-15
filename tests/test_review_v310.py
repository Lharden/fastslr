"""Regressions from the 3.1.0 merge review (pre-IP audit + 3.0.1).

- Block names in the terms table ignore case: 'Global' is the GLOBAL (T0) block,
  and 'Ctx' / 'CTX' are one block whose terms all survive.
- Boolean and integer parameters read from JSON strings parse to what they say.
- The xlsx export never turns corpus text into a live formula.
- doctor and run agree on a corpus without a title or abstract column.
- The F08 preflight refuses blank and duplicate IDs in run, preview and coverage.
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest

from fastslr.app.controller import (
    analyze_coverage,
    inspect_run_setup,
    preview_triage,
    run_triage,
)
from fastslr.core.config import load_global_params, parse_terms_csv
from fastslr.core.io import export_results
from fastslr.core.normalization import NormalizationEngine


def _write_terms_csv(path: Path, rows: list[str]) -> Path:
    header = "block,kind,term,level,section_scope,is_regex"
    path.write_text(header + "\n" + "\n".join(rows) + "\n", encoding="utf-8")
    return path


# ── block names ignore case ──────────────────────────────────────────────────


@pytest.mark.parametrize("spelling", ["Global", "global"])
def test_global_block_in_any_case_is_the_t0_block(tmp_path: Path, spelling: str) -> None:
    csv = _write_terms_csv(
        tmp_path / "terms.csv",
        [f"{spelling},flag,case studies,,any,0", "T1A,pos,oil and gas,1,any,0"],
    )
    base_config = {"global": {"LIMITES_APROVADO": {"1": 10}}}

    config = parse_terms_csv(csv, base_config)

    assert config["_domain_blocks"] == ["T1A"]
    assert [e["term"] for e in config["T0"]["anti"]["flag"]] == ["case studies"]
    assert config["global"]["LIMITES_APROVADO"] == {"1": 10}
    assert any("ignore case" in w and spelling in w for w in config["_parse_warnings"])


def test_blocks_that_differ_only_in_case_keep_all_their_terms(tmp_path: Path) -> None:
    csv = _write_terms_csv(
        tmp_path / "terms.csv",
        ["Ctx,pos,oil,1,any,0", "CTX,pos,gas,1,any,0"],
    )

    config = parse_terms_csv(csv, {"global": {"BLOCK_ORDER": ["CTX"]}})

    assert sorted(e["term"] for e in config["CTX"]["positives"]) == ["gas", "oil"]
    assert any("'CTX'" in w and "ignore case" in w for w in config["_parse_warnings"])


def test_canonical_block_names_produce_no_warning(tmp_path: Path) -> None:
    csv = _write_terms_csv(
        tmp_path / "terms.csv",
        ["GLOBAL,flag,case studies,,any,0", "T1A,pos,oil,1,any,0", "T1A,pos,gas,1,any,0"],
    )

    config = parse_terms_csv(csv, {"global": {}})

    assert config["_parse_warnings"] == []


# ── JSON strings in parameters ───────────────────────────────────────────────


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        (True, True),
        (False, False),
        ("false", False),
        ("0", False),
        ("no", False),
        ("true", True),
        ("1", True),
        ("yes", True),
        (0, False),
        (1, True),
    ],
)
def test_boolean_parameters_parse_strings(raw: object, expected: bool) -> None:
    params = load_global_params(
        {
            "ENABLE_PROXIMITY_DETECTION": raw,
            "FAIL_FAST_GLOBAL": raw,
            "ENABLE_SPECIAL_APPROVAL_RULE": raw,
            "REQUIRE_NON_WEAK_TERM_FOR_APPROVAL": raw,
        }
    )

    assert params.enable_proximity_detection is expected
    assert params.fail_fast_enabled is expected
    assert params.enable_special_approval_rule is expected
    assert params.require_non_weak_term_for_approval is expected


@pytest.mark.parametrize("raw", ["maybe", "", 2, None])
def test_unreadable_boolean_parameter_is_refused(raw: object) -> None:
    with pytest.raises(ValueError, match="FAIL_FAST_GLOBAL"):
        load_global_params({"FAIL_FAST_GLOBAL": raw})


def test_min_approved_blocks_is_an_integer_of_at_least_one() -> None:
    assert load_global_params({"MIN_APPROVED_BLOCKS": "2"}).min_approved_blocks == 2
    assert load_global_params({"MIN_APPROVED_BLOCKS": None}).min_approved_blocks is None
    with pytest.raises(ValueError, match="MIN_APPROVED_BLOCKS"):
        load_global_params({"MIN_APPROVED_BLOCKS": 0})
    with pytest.raises(ValueError, match="MIN_APPROVED_BLOCKS"):
        load_global_params({"MIN_APPROVED_BLOCKS": "two"})


def test_token_unit_outside_the_allowlist_is_a_visible_warning() -> None:
    from fastslr.app.controller import validate_config

    def warnings(token_unit: str) -> list[str]:
        config = {"global": {"TOKEN_UNIT_FOR_GAPS": token_unit}}
        issues = validate_config(config, require_resolved_blocks=False)
        return [i.message for i in issues if "TOKEN_UNIT_FOR_GAPS" in i.message]

    assert len(warnings("[A-Za-z]+")) == 1
    assert warnings(r"\S+") == []
    assert warnings(r"\w+") == []


# ── xlsx export keeps text as text ───────────────────────────────────────────


@pytest.mark.parametrize("engine", ["openpyxl", "xlsxwriter"])
def test_xlsx_export_writes_formula_like_text_as_text(tmp_path: Path, engine: str) -> None:
    from openpyxl import load_workbook

    hostile = '=HYPERLINK("http://example.invalid/","click")'
    df = pd.DataFrame({"ID": ["A1", "A2"], "Title": [hostile, "Plain title"]})
    output = {"xlsx": True, "csv": False, "xlsx_engine": engine}

    exported = export_results(df, tmp_path / "r", {"output": output})

    sheet = load_workbook(exported["xlsx"]).active
    assert sheet is not None
    cells = {
        sheet.cell(row=r, column=2).value: sheet.cell(row=r, column=2).data_type for r in (2, 3)
    }
    assert cells == {hostile: "s", "Plain title": "s"}


# ── doctor and run agree; F08 preflight in every execution flow ──────────────


def _config(tmp_path: Path) -> Path:
    path = tmp_path / "config.json"
    path.write_text(
        json.dumps(
            {
                "global": {"BLOCK_ORDER": ["T1A"]},
                "T1A": {"positives": [{"term": "oil", "level": 1}]},
                "fields": {"id": "key", "title": "title", "abstract": "abstract"},
            }
        ),
        encoding="utf-8",
    )
    return path


def _articles(tmp_path: Path, columns: dict[str, list[str]]) -> Path:
    path = tmp_path / "articles.csv"
    pd.DataFrame(columns).to_csv(path, sep=";", index=False)
    return path


def test_doctor_reports_a_missing_abstract_column_as_the_error_run_raises(tmp_path: Path) -> None:
    config_path = _config(tmp_path)
    input_path = _articles(tmp_path, {"key": ["A1"], "title": ["Oil fields"], "year": ["2020"]})

    inspection = inspect_run_setup(input_path=input_path, config_path=config_path)
    with pytest.raises(ValueError, match="abstract"):
        run_triage(input_path, config_path, output_dir=tmp_path / "out")

    assert any("abstract" in error for error in inspection.errors)


@pytest.mark.parametrize("flow", ["run", "preview", "coverage"])
@pytest.mark.parametrize("keys", [["A1", "A1"], ["A1", ""]], ids=["duplicate", "blank"])
def test_f08_preflight_refuses_bad_ids_in_every_flow(
    tmp_path: Path, flow: str, keys: list[str]
) -> None:
    config_path = _config(tmp_path)
    input_path = _articles(
        tmp_path, {"key": keys, "title": ["Oil", "Gas"], "abstract": ["Oil", "Gas"]}
    )
    out = tmp_path / "out"

    with pytest.raises(ValueError, match="Invalid record IDs"):
        if flow == "run":
            run_triage(input_path, config_path, output_dir=out)
        elif flow == "preview":
            preview_triage(input_path, config_path)
        else:
            analyze_coverage(input_path, config_path)

    assert not out.exists()


# ── normalization and matching ───────────────────────────────────────────────


def _engine(**rules: dict) -> NormalizationEngine:
    return NormalizationEngine({"enabled": True, **rules})


def test_nfkc_gives_one_form_to_text_that_only_looks_the_same() -> None:
    import unicodedata

    engine = NormalizationEngine({"enabled": False})
    nfd = unicodedata.normalize("NFD", "Gestão de cadeia")

    assert engine.normalize(nfd) == engine.normalize("Gestão de cadeia") == "gestão de cadeia"
    assert engine.normalize("ＡＩ systems") == "ai systems"
    assert _engine(abbreviations={"AI": "artificial intelligence"}).normalize("ＡＩ systems") == (
        "artificial intelligence systems"
    )


def test_lowercase_abbreviation_key_matches_any_case() -> None:
    engine = _engine(abbreviations={"iot": "internet of things"})

    assert engine.normalize("IoT platforms and IOT and iot") == (
        "internet of things platforms and internet of things and internet of things"
    )


def test_abbreviation_key_with_uppercase_keeps_its_case() -> None:
    engine = _engine(abbreviations={"ML": "machine learning"})

    assert engine.normalize("5 ml of ML") == "5 ml of machine learning"


def test_keys_with_non_word_edges_expand() -> None:
    engine = _engine(
        abbreviations={"C#": "csharp", ".NET": "dotnet"},
        compound_variants={"c++": "cplusplus"},
    )

    assert engine.normalize("using C# and .NET with c++ code") == (
        "using csharp and dotnet with cplusplus code"
    )
    assert engine.normalize("ABC#D") == "abc#d"


def test_abbreviation_rules_that_differ_only_in_case_are_reported(tmp_path: Path) -> None:
    path = tmp_path / "terms.csv"
    path.write_text(
        "block,kind,term,level,section_scope,is_regex,normalization_type,normalization_target\n"
        ",,ML,,,0,abbreviation,machine learning\n"
        ",,ml,,,0,abbreviation,millilitre\n"
        "T1A,pos,oil,1,any,0,,\n",
        encoding="utf-8",
    )

    config = parse_terms_csv(path, {"global": {}})

    assert any("differ only in case" in w for w in config["_parse_warnings"])


@pytest.mark.parametrize("term", ["*", "**", "?*"])
def test_wildcard_only_term_is_skipped(tmp_path: Path, term: str) -> None:
    csv = _write_terms_csv(
        tmp_path / "terms.csv", [f"T1A,pos,{term},1,any,0", "T1A,pos,oil,1,any,0"]
    )

    config = parse_terms_csv(csv, {"global": {}})

    assert [e["term"] for e in config["T1A"]["positives"]] == ["oil"]
    assert any("only of wildcards" in w for w in config["_parse_warnings"])
