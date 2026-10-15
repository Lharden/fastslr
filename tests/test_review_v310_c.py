"""Regressions from the second verification round of the 3.1.0 review.

- A control character in a title never makes the xlsx export fail.
- The CSV is the raw data (decision of the user), and the xlsx never holds a formula.
- The protocol records the booleans the engine ran with.
- A block is spelled as its first accepted row; T0 is reserved; an infinite level is invalid.
- A parse error is reported as such, not as a narrow table.
- The highlight reads the same NFKC form as scoring, and marks the original text.
- Rule keys go through NFKC like the text.
- ERROR_POLICY and MIN_APPROVED_BLOCKS are checked by validate.
- The CLI keeps the traceback of bugs (NotImplementedError, RecursionError).
- diff refuses repeated blank IDs, which the merge would pair with each other.
- An old protocol migrated to 2.2 gets the 2.2 meaning, not only the label.
- Symbol keys go through NFKC instead of being dropped.
- ERROR_POLICY is refused where it is read.
- MIN_APPROVED_BLOCKS is validated as the run reads it.
"""

from __future__ import annotations

import json
import re
import unicodedata
import zipfile
from pathlib import Path

import pandas as pd
import pytest

import fastslr.app.cli as cli_module
from fastslr.app.controller import diff_results, run_triage, validate_config
from fastslr.core.config import parse_terms_csv
from fastslr.core.io import export_results, highlight_text, load_table_safe
from fastslr.core.normalization import NormalizationEngine


def _terms(tmp_path: Path, rows: list[str]) -> Path:
    path = tmp_path / "terms.csv"
    path.write_text(
        "block,kind,term,level,section_scope,is_regex\n" + "\n".join(rows) + "\n",
        encoding="utf-8",
    )
    return path


# ── export ───────────────────────────────────────────────────────────────────


@pytest.mark.parametrize("engine", ["openpyxl", "xlsxwriter"])
def test_control_characters_do_not_break_the_xlsx_export(tmp_path: Path, engine: str) -> None:
    from openpyxl import load_workbook

    df = pd.DataFrame({"ID": ["a"], "Title_Highlighted": ["foo\x0bbar\x0cbaz"], "Score": [1.5]})
    output = {"xlsx": True, "csv": False, "xlsx_engine": engine}

    exported = export_results(df, tmp_path / "r", {"output": output})

    sheet = load_workbook(exported["xlsx"]).active
    assert sheet is not None
    assert sheet.cell(row=2, column=2).value == "foo bar baz"


def test_csv_is_the_raw_data_cell_for_cell(tmp_path: Path) -> None:
    # The user's decision for 3.1.0: the CSV stays raw (no leading apostrophe), so
    # it round-trips and diffs; the xlsx is the file that never holds a formula.
    titles = ['=HYPERLINK("x")', "+1+1", "@SUM(1)", "-2+3", "Plain title"]
    df = pd.DataFrame({"ID": list("abcde"), "Title": titles, "Score": [-1.5, 2, 3, 4, 5]})
    output = {"xlsx": False, "csv": True, "csv_sep": ";", "csv_decimal": "."}

    exported = export_results(df, tmp_path / "r", {"output": output, "encoding": "utf-8"})

    back = pd.read_csv(exported["csv"], sep=";", dtype=str, keep_default_na=False)
    assert list(back["Title"]) == titles
    assert back["Score"].iloc[0] == "-1.50"


def test_protocol_records_the_booleans_the_engine_ran_with(tmp_path: Path) -> None:
    config = {
        "global": {"BLOCK_ORDER": ["T1A"], "FAIL_FAST_GLOBAL": "false"},
        "T1A": {"positives": [{"term": "oil", "level": 1}]},
        "fields": {"id": "key", "title": "title", "abstract": "abstract"},
    }
    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps(config), encoding="utf-8")
    input_path = tmp_path / "in.csv"
    pd.DataFrame({"key": ["A1"], "title": ["Oil"], "abstract": ["Oil"]}).to_csv(
        input_path, sep=";", index=False
    )

    result = run_triage(input_path, config_path, output_dir=tmp_path / "out")

    protocol = json.loads((tmp_path / "out" / "protocol.json").read_text(encoding="utf-8"))
    assert protocol["protocol_version"] == "2.2"
    assert protocol["configuration"]["fail_fast"] is False
    with zipfile.ZipFile(result.artifact_paths["academic_package"]) as zf:
        assert json.loads(zf.read("appendix_manifest.json"))["manifest_version"] == 2


# ── terms table ──────────────────────────────────────────────────────────────


def test_block_is_spelled_as_its_first_accepted_row(tmp_path: Path) -> None:
    csv = _terms(tmp_path, ["ctx,badkind,oil,1,any,0", "CTX,pos,gas,1,any,0"])

    config = parse_terms_csv(csv, {"global": {}})

    assert config["_domain_blocks"] == ["CTX"]
    assert not any("ignore case" in w for w in config["_parse_warnings"])


def test_t0_is_reserved_for_the_global_block(tmp_path: Path) -> None:
    csv = _terms(
        tmp_path,
        ["T0,pos,oil,1,any,0", "GLOBAL,flag,case studies,,any,0", "T1A,pos,gas,1,any,0"],
    )

    config = parse_terms_csv(csv, {"global": {}})

    assert config["_domain_blocks"] == ["T1A"]
    assert [e["term"] for e in config["T0"]["anti"]["flag"]] == ["case studies"]
    assert config["T0"]["positives"] == []
    assert any("reserved" in w for w in config["_parse_warnings"])


@pytest.mark.parametrize("level", ["inf", "1e400", "nan"])
def test_level_that_is_not_a_whole_number_is_invalid_not_a_crash(
    tmp_path: Path, level: str
) -> None:
    csv = _terms(tmp_path, [f"T1A,pos,oil,{level},any,0", "T1A,pos,gas,1,any,0"])

    config = parse_terms_csv(csv, {"global": {}})

    oil = next(e for e in config["T1A"]["positives"] if e["term"] == "oil")
    assert oil["level"] is None


def test_parse_error_is_reported_with_the_column_count(tmp_path: Path) -> None:
    path = tmp_path / "bad.csv"
    path.write_text("ID,Title,Abstract\n1,a,b\n2,c,d,e\n3,f,g\n", encoding="utf-8")

    with pytest.raises(ValueError) as caught:
        load_table_safe(path)

    assert "also failed to parse" in str(caught.value)
    assert caught.value.__cause__ is not None


def test_narrow_table_keeps_its_message_when_a_wrong_separator_fails(tmp_path: Path) -> None:
    path = tmp_path / "narrow.csv"
    path.write_text('ID,Title\n1,"a;b;c"\n2,d\n', encoding="utf-8")

    with pytest.raises(ValueError, match=r"at most 2 column\(s\)"):
        load_table_safe(path)


# ── highlight and NFKC ───────────────────────────────────────────────────────


def test_highlight_marks_a_term_that_scored_on_nfkc_text() -> None:
    nfd = unicodedata.normalize("NFD", "modelo de gestão de projetos")
    term = {"pattern": re.compile(r"\bgestão\b", re.IGNORECASE), "scope": "any"}

    marked = highlight_text(nfd, [term], "abstract")

    assert marked.startswith("modelo de ***")
    assert marked.endswith("*** de projetos")
    assert unicodedata.normalize("NFC", marked.split("***")[1]) == "GESTÃO"


def test_highlight_maps_full_width_text_back_to_the_original() -> None:
    term = {"pattern": re.compile(r"\bai\b", re.IGNORECASE), "scope": "any"}

    assert highlight_text("ＡＩ systems", [term], "title") == "***ＡＩ*** systems"


def test_highlight_of_nfkc_text_is_unchanged() -> None:
    term = {"pattern": re.compile(r"\boil\b", re.IGNORECASE), "scope": "any"}

    assert highlight_text("Oil and gas", [term], "title") == "***OIL*** and gas"


def test_rule_keys_go_through_nfkc() -> None:
    engine = NormalizationEngine({"enabled": True, "abbreviations": {"ＭＬ": "machine learning"}})

    assert engine.normalize("ML models") == "machine learning models"


# ── validation and CLI ───────────────────────────────────────────────────────


def test_unknown_error_policy_is_an_error() -> None:
    issues = validate_config({"global": {"ERROR_POLICY": "skip"}}, require_resolved_blocks=False)

    assert any(i.level == "error" and "ERROR_POLICY" in i.message for i in issues)


def test_min_approved_blocks_above_the_block_count_is_a_warning() -> None:
    config = {
        "global": {"BLOCK_ORDER": ["T1A"], "DECISION_POLICY": "k_of_n", "MIN_APPROVED_BLOCKS": 3},
        "T1A": {"positives": [{"term": "oil"}]},
    }

    issues = validate_config(config)

    assert any(i.level == "warning" and "MIN_APPROVED_BLOCKS" in i.message for i in issues)


@pytest.mark.parametrize("error", [NotImplementedError, RecursionError])
def test_cli_keeps_the_traceback_of_bugs(
    monkeypatch: pytest.MonkeyPatch, error: type[RuntimeError]
) -> None:
    def broken() -> None:
        raise error("bug")

    monkeypatch.setattr(cli_module, "app", broken)

    with pytest.raises(error):
        cli_module.main()


def test_diff_refuses_repeated_blank_ids(tmp_path: Path) -> None:
    # The outer merge pairs '' with '': repeated blank IDs would invent transitions,
    # so they are refused like any repeated ID (third verification round, 2 highs).
    a = tmp_path / "a.csv"
    b = tmp_path / "b.csv"
    pd.DataFrame(
        {"ID": ["1", "", ""], "Final_Decision": ["APPROVED_FINAL"] + ["REJECTED_FINAL"] * 2}
    ).to_csv(a, sep=";", index=False)
    pd.DataFrame(
        {
            "ID": ["1", "", ""],
            "Final_Decision": ["APPROVED_FINAL", "APPROVED_FINAL", "REJECTED_FINAL"],
        }
    ).to_csv(b, sep=";", index=False)

    with pytest.raises(ValueError, match=r"repeated IDs.*\(blank\)"):
        diff_results(a, b)


# ── stabilization of the third round ─────────────────────────────────────────


def test_old_protocol_migrated_to_2_2_gets_the_2_2_meaning() -> None:
    from fastslr.core.io import migrate_protocol_snapshot

    old = {
        "protocol_version": "2.1",
        "execution_id": "run_x",
        "inputs": {
            "input_file": "C:/Users/alice/secret/input.csv",
            "terms_file": "C:\\Users\\alice\\terms.csv",
        },
        "configuration": {"fail_fast": "false", "enable_special_approval": "1"},
        "artifacts": {"results_path": "/home/alice/out/r.xlsx"},
    }

    migrated = migrate_protocol_snapshot(old)

    assert migrated["protocol_version"] == "2.2"
    assert migrated["inputs"]["input_file"] == "input.csv"
    assert migrated["inputs"]["terms_file"] == "terms.csv"
    assert migrated["artifacts"]["results_path"] == "r.xlsx"
    assert migrated["configuration"]["fail_fast"] is False
    assert migrated["configuration"]["enable_special_approval"] is True


def test_symbol_key_that_nfkc_rewrites_still_applies() -> None:
    engine = NormalizationEngine({"enabled": True, "symbol_replacements": {"½": "half"}})

    assert "half" in engine.normalize("a ½ dose")


def test_unknown_error_policy_is_refused_where_it_is_read() -> None:
    from fastslr.core.config import load_global_params

    with pytest.raises(ValueError, match="ERROR_POLICY"):
        load_global_params({"ERROR_POLICY": "FLAG"})


@pytest.mark.parametrize("raw", ["5", 5.0])
def test_min_approved_blocks_is_validated_as_the_run_reads_it(raw: object) -> None:
    config = {
        "global": {"BLOCK_ORDER": ["T1A"], "DECISION_POLICY": "k_of_n", "MIN_APPROVED_BLOCKS": raw},
        "T1A": {"positives": [{"term": "oil"}]},
    }

    issues = validate_config(config)

    assert any(i.level == "warning" and "MIN_APPROVED_BLOCKS" in i.message for i in issues)


def test_empty_match_marks_nothing_on_the_direct_path() -> None:
    # 'x*' matches the empty string at every position of 'Oil', away from any real match.
    term = {"pattern": re.compile(r"x*", re.IGNORECASE), "scope": "any"}

    assert highlight_text("Oil", [term], "title") == "Oil"


def test_xlsx_cleans_noncharacters_and_column_names(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    df = pd.DataFrame({"Ti\x0btle": ["a\ufffeb"]})

    exported = export_results(df, tmp_path / "r", {"output": {"xlsx": True, "csv": False}})

    sheet = load_workbook(exported["xlsx"]).active
    assert sheet is not None
    assert (sheet.cell(row=1, column=1).value, sheet.cell(row=2, column=1).value) == (
        "Ti tle",
        "a b",
    )
